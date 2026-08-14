"""Tests for export functions."""

import json
import tempfile
from pathlib import Path

import pytest
from Bio.Seq import Seq

from knockknock.engine import (
    AnalysisParams,
    make_shrnas,
    process_shrnas,
)
from knockknock.export import (
    no_candidates_message,
    resolve_export_path,
    to_csv_text,
    to_fasta,
    to_json_text,
    to_xlsx,
    write_candidates_export,
)


@pytest.fixture
def sample_results():
    mrnas = [("test", Seq("ATGCGTAGCTAGCTAGCTAGCTGAT"))]
    params = AnalysisParams()
    candidates = make_shrnas(mrnas, params)
    results, stats = process_shrnas(candidates, mrnas, params)
    return results, stats, params


def test_no_candidates_message_names_the_short_input_case():
    """Zero windows scanned names the short-input cause specifically, rather
    than falling through to the generic mixed-rejection wording."""
    stats = {"windows_scanned": 0, "unique_guides": 0}
    message = no_candidates_message(stats, AnalysisParams())
    assert "shorter than the guide length" in message


def test_no_candidates_message_names_the_all_wildcard_case():
    """Every unique guide overlapping a wildcard base gets its own message,
    naming the guide length so the reader knows why nothing survived."""
    stats = {"windows_scanned": 9, "unique_guides": 9, "wildcard_filtered": 9}
    message = no_candidates_message(stats, AnalysisParams())
    assert "overlap a wildcard base" in message
    assert "22 nt guide length" in message


def test_no_candidates_message_breaks_down_a_mixed_rejection():
    """A rejection spread across several filters lists each one with its
    count, and leaves out any filter that rejected nothing."""
    stats = {
        "windows_scanned": 10,
        "unique_guides": 10,
        "wildcard_filtered": 4,
        "gc_filtered": 3,
        "motif_filtered": 3,
    }
    message = no_candidates_message(stats, AnalysisParams(gc_min=30, gc_max=60))
    assert "4 overlapped a wildcard base" in message
    assert "3 fell outside GC 30–60%" in message
    assert "3 hit an avoided motif" in message
    # Filters that rejected nothing are left out rather than listed as zero.
    assert "scored below" not in message


def test_no_candidates_message_avoids_cli_flag_names():
    """The TUI renders this too, and has no flags to type."""
    stats = {"windows_scanned": 5, "unique_guides": 5, "gc_filtered": 5}
    assert "--" not in no_candidates_message(stats, AnalysisParams())


def test_to_csv(sample_results):
    """`to_csv_text` written to disk has the expected data columns and one
    row per candidate, plus a one-line header."""
    results, stats, params = sample_results
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False) as f:
        path = Path(f.name)
    try:
        path.write_text(to_csv_text(results, params), encoding="utf-8")
        content = path.read_text()
        assert "Guide_Sequence" in content
        assert "Efficacy_Score" in content
        non_meta = [ln for ln in content.splitlines() if not ln.startswith("#")]
        # 1 header + N rows
        assert len(non_meta) == len(results) + 1
    finally:
        path.unlink(missing_ok=True)


def test_csv_has_metadata_header(sample_results):
    """The CSV leads with `#`-prefixed provenance lines (version, timestamp)
    ahead of the data rows, so a result file is reproducible on its own."""
    results, _, params = sample_results
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False) as f:
        path = Path(f.name)
    try:
        path.write_text(to_csv_text(results, params), encoding="utf-8")
        lines = path.read_text().splitlines()
        assert lines[0].startswith("# knockknock_version: ")
        assert any(ln.startswith("# timestamp_utc: ") for ln in lines[:20])
    finally:
        path.unlink(missing_ok=True)


def test_to_json(sample_results):
    """JSON export has top-level `meta` and `candidates` keys, one candidate
    object per result, and the expected fields on each."""
    results, stats, params = sample_results
    with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as f:
        path = Path(f.name)
    try:
        path.write_text(to_json_text(results, params), encoding="utf-8")
        data = json.loads(path.read_text())
        assert "meta" in data and "candidates" in data
        assert len(data["candidates"]) == len(results)
        assert "guide_sequence" in data["candidates"][0]
        assert "efficacy_score" in data["candidates"][0]
        assert data["meta"]["guide_length"] == params.length
    finally:
        path.unlink(missing_ok=True)


def test_to_xlsx(sample_results):
    """`to_xlsx` writes a non-empty workbook to the given path."""
    results, stats, params = sample_results
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = Path(f.name)
    try:
        to_xlsx(results, params, path)
        assert path.exists()
        assert path.stat().st_size > 0
    finally:
        path.unlink(missing_ok=True)


def test_to_xlsx_writes_empty_result_workbook(tmp_path: Path):
    """Zero candidates still produce a real workbook — a header-only sheet
    the user can open, not a failed or truncated export."""
    path = tmp_path / "empty.xlsx"
    to_xlsx([], AnalysisParams(), path)
    assert path.exists() and path.stat().st_size > 0
    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert wb["shRNA Candidates"].max_row == 1


def test_csv_text_matches_file(sample_results):
    """The string `to_csv_text` returns matches what actually lands on disk
    (comparing everything but the timestamped metadata lines, which can
    drift by a second between the two calls)."""
    results, _, params = sample_results
    text = to_csv_text(results, params)
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False) as f:
        path = Path(f.name)
    try:
        path.write_text(to_csv_text(results, params), encoding="utf-8")
        # Timestamps in the metadata header can drift by a second between
        # the two calls, so compare everything but the `# ` comment lines.
        rows_from_text = [ln for ln in text.splitlines() if not ln.startswith("#")]
        rows_from_file = [
            ln for ln in path.read_text().splitlines() if not ln.startswith("#")
        ]
        assert rows_from_text == rows_from_file
    finally:
        path.unlink(missing_ok=True)


def test_json_text_matches_file(sample_results):
    """The string `to_json_text` returns matches what actually lands on
    disk."""
    results, _, params = sample_results
    text = to_json_text(results, params)
    with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as f:
        path = Path(f.name)
    try:
        path.write_text(to_json_text(results, params), encoding="utf-8")
        assert (
            json.loads(text)["candidates"] == json.loads(path.read_text())["candidates"]
        )
    finally:
        path.unlink(missing_ok=True)


def test_to_fasta(sample_results):
    """FASTA output is one `>`-header line plus one sequence line per
    candidate, and the sequence line is the guide's antisense strand."""
    results, _, params = sample_results
    text = to_fasta(results, params)
    lines = text.splitlines()
    assert len(lines) == 2 * len(results)
    assert lines[0].startswith(">")
    assert lines[1] == results[0].antisense


def test_to_fasta_empty():
    """Zero candidates produce an empty string, not a dangling header or
    trailing newline."""
    assert to_fasta([], AnalysisParams()) == ""


@pytest.mark.parametrize("fmt,ext", [("fasta", ".fasta"), ("csv", ".csv"), ("json", ".json")])
def test_write_candidates_export_matches_the_matching_to_function(
    sample_results, tmp_path, fmt, ext
):
    """`write_candidates_export` is the single dispatch both the explorer and
    oligo-designer TUI screens call for their candidate export; each format
    branch must write exactly what its own `to_*` function would return, so
    the two screens can't quietly drift apart on a shared format."""
    results, _, params = sample_results
    path = tmp_path / f"out{ext}"
    write_candidates_export(fmt, results, params, path)

    if fmt == "fasta":
        assert path.read_text() == to_fasta(results, params)
    elif fmt == "csv":
        # Timestamps in the metadata header can drift by a second between
        # the two calls, so compare everything but the `# ` comment lines.
        written = [ln for ln in path.read_text().splitlines() if not ln.startswith("#")]
        expected = [
            ln for ln in to_csv_text(results, params).splitlines() if not ln.startswith("#")
        ]
        assert written == expected
    else:  # json
        assert json.loads(path.read_text())["candidates"] == json.loads(
            to_json_text(results, params)
        )["candidates"]


def test_write_candidates_export_xlsx_writes_to_the_given_path(sample_results, tmp_path):
    """The `xlsx` branch is the one format that needs a path instead of
    returning text (xlsxwriter can't render to a string) — it must write
    straight to `out_path` rather than returning anything."""
    results, _, params = sample_results
    path = tmp_path / "out.xlsx"
    assert write_candidates_export("xlsx", results, params, path) is None
    assert path.exists() and path.stat().st_size > 0


def test_rna_output(sample_results):
    """`rna_output=True` converts the exported guide_sequence to RNA (U, no
    T) via `maybe_rna`'s per-field conversion, not a blanket text replace
    over the rendered output."""
    results, stats, params = sample_results
    params_rna = AnalysisParams(rna_output=True)
    with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as f:
        path = Path(f.name)
    try:
        path.write_text(to_json_text(results, params_rna), encoding="utf-8")
        data = json.loads(path.read_text())
        first = data["candidates"][0]
        assert "U" in first["guide_sequence"] or "T" not in first["guide_sequence"]
    finally:
        path.unlink(missing_ok=True)


def test_xlsx_has_run_info_sheet(sample_results):
    """The workbook carries a second "Run Info" sheet of key/value
    provenance pairs (tool version, guide length, ...) alongside the
    candidate table."""
    from openpyxl import load_workbook

    results, _, params = sample_results
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = Path(f.name)
    try:
        to_xlsx(results, params, path)
        wb = load_workbook(path)
        assert "Run Info" in wb.sheetnames
        ws = wb["Run Info"]
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0] == ("key", "value")
        keys = [r[0] for r in rows[1:]]
        assert "knockknock_version" in keys
        assert "guide_length" in keys
    finally:
        path.unlink(missing_ok=True)


def test_xlsx_uses_reference_bounds(sample_results):
    """The Efficacy Score colour scale's min/max should equal bounds derived
    from the reference coefficient table (regression for the
    stale-hardcoded bounds)."""
    from openpyxl import load_workbook

    from knockknock.reference import load_reference_data

    results, stats, params = sample_results
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = Path(f.name)
    try:
        to_xlsx(results, params, path)
        ref = load_reference_data()
        expected_min = float(ref.coeffs[["A", "C", "G", "T"]].min(axis=1).sum())
        expected_max = float(ref.coeffs[["A", "C", "G", "T"]].max(axis=1).sum())

        wb = load_workbook(path)
        ws = wb.active
        cf_rules = list(ws.conditional_formatting._cf_rules.items())
        score_rule = None
        for _rng, rules in cf_rules:
            for rule in rules:
                if rule.type == "colorScale" and len(rule.colorScale.cfvo) >= 3:
                    cfvo = rule.colorScale.cfvo
                    # Match the score range (compare against expected min)
                    if abs(float(cfvo[0].val) - expected_min) < 0.01:
                        score_rule = rule
                        break
        assert score_rule is not None, "score colour-scale rule not found"
        assert abs(float(score_rule.colorScale.cfvo[0].val) - expected_min) < 0.01
        assert abs(float(score_rule.colorScale.cfvo[2].val) - expected_max) < 0.01
    finally:
        path.unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# resolve_export_path: with no explicit -o, the file lands beside the input
# file. Without a real input file to sit beside (stdin, a paste, or no
# reference_path at all), it falls back to the current working directory.
# ---------------------------------------------------------------------------


def test_resolve_export_path_uses_cwd_when_no_reference_path(tmp_path):
    """With nothing to name the file after, the export still lands in cwd."""
    path = resolve_export_path(".csv", reference_path=None)
    assert path.parent == tmp_path.resolve()
    assert path.suffix == ".csv"


def test_resolve_export_path_uses_cwd_for_stdin_input(tmp_path):
    """'-' (piped stdin) is not a path to sit beside; cwd is used."""
    path = resolve_export_path(".csv", reference_path="-")
    assert path.parent == tmp_path.resolve()


def test_resolve_export_path_uses_the_input_directory(tmp_path):
    """The export lands beside the input FASTA, not in cwd, so the user finds
    it next to the file they just ran the command on."""
    elsewhere = tmp_path / "reference_data"
    elsewhere.mkdir()
    fasta = elsewhere / "targets.fa"
    fasta.write_text(">x\nACGT\n")

    path = resolve_export_path(".csv", reference_path=str(fasta))

    assert path.parent == elsewhere.resolve()
    assert path.parent != tmp_path.resolve()


def test_resolve_export_path_names_the_file_after_the_input_stem(tmp_path):
    """The input's stem still names the export, so it stays identifiable."""
    fasta = tmp_path / "TP53_isoforms.fa"
    fasta.write_text(">x\nACGT\n")

    path = resolve_export_path(".csv", reference_path=str(fasta))

    assert path.name.startswith("TP53_isoforms_")


def test_resolve_export_path_honours_an_explicit_output_path(tmp_path):
    """An explicit -o always wins over the cwd default."""
    target = tmp_path / "somewhere" / "mine"
    target.parent.mkdir()

    path = resolve_export_path(".csv", user_output=str(target))

    assert path == target.with_suffix(".csv").resolve()


def test_resolve_export_path_label_separates_two_exports_in_the_same_second(tmp_path):
    """Candidate and oligo exports share extensions; the label keeps them from
    resolving to the same filename and silently overwriting each other."""
    candidates = resolve_export_path(".csv", reference_path=None, label="candidates")
    oligos = resolve_export_path(".csv", reference_path=None, label="oligos")
    assert candidates != oligos


def test_resolve_export_path_honors_user_output_verbatim(tmp_path):
    """An explicit user_output is used as given (only the extension is
    normalized), independent of reference_path."""
    target = tmp_path / "custom_name"
    path = resolve_export_path(".csv", user_output=str(target), reference_path=None)
    assert path == target.with_suffix(".csv").resolve()
